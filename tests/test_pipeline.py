"""テストはすべて架空のデータで行う（実データをリポジトリに置かないこと）。"""
import tempfile
import unittest
from datetime import date
from pathlib import Path

import pandas as pd

from takarazuka_data.classify import ANNUAL, GOODS, LARGE, SMALL, SOCIAL, TEA_PARTY, UNKNOWN, classify_sheet
from takarazuka_data.classify import APPLICATION, ROSTER
from takarazuka_data.columns import (
    clean_amount, clean_phone, clean_postal, clean_prefecture, clean_seat, map_header, parse_show)
from takarazuka_data.pipeline import run, select_events, write_outputs

STD = ["name(漢字)", "name katakana", "Family name", "First name", "mail", "phone-number",
       "post-number", "prefecture", "address 1", "address 2"]


class ClassifyTest(unittest.TestCase):
    def test_categories(self):
        cases = {
            "2026_0121_大公演_渤海": (LARGE, "渤海", date(2026, 1, 21)),
            "2026_0415_small_春の夢": (SMALL, "春の夢", date(2026, 4, 15)),
            "渤海_お茶会": (TEA_PARTY, "渤海", None),
            "春の夢_親睦会": (SOCIAL, "春の夢", None),
            "渤海_グッズ": (GOODS, "渤海", None),
            "2026_1201_クリスマス会": (ANNUAL, "クリスマス会", date(2026, 12, 1)),
            "２０２６＿０１２１＿大公演＿渤海": (LARGE, "渤海", date(2026, 1, 21)),  # 全角
        }
        for sheet, (cat, name, d) in cases.items():
            info = classify_sheet(sheet)
            self.assertEqual((info.category, info.name, info.event_date), (cat, name, d), sheet)

    def test_unknown_and_bad_date(self):
        self.assertEqual(classify_sheet("Sheet1").category, UNKNOWN)
        self.assertIsNotNone(classify_sheet("2026_1340_大公演_渤海").error)


class ColumnsTest(unittest.TestCase):
    def test_headers(self):
        self.assertEqual(map_header("name(漢字)"), "name_kanji")
        self.assertEqual(map_header("フリガナ"), "name_kana")
        self.assertEqual(map_header("phone-number"), "phone")
        self.assertEqual(map_header("Ｅｎｔｒｙ＿ｆｅｅ"), "entry_fee")
        self.assertEqual(map_header("住所①"), "address1")

    def test_values(self):
        self.assertEqual(clean_phone(9012345678)[0], "09012345678")
        self.assertEqual(clean_phone("０９０－１２３４－５６７８")[0], "09012345678")
        self.assertEqual(clean_postal(600000)[0], "060-0000")
        self.assertEqual(clean_postal("〒１２３ー４５６７")[0], "123-4567")
        self.assertEqual(clean_prefecture("東京"), "東京都")
        self.assertEqual(clean_amount("¥12,000円")[0], 12000)
        self.assertIsNotNone(clean_amount("abc")[1])

    def test_ticket_values(self):
        self.assertEqual(clean_seat("ｓｓ席", LARGE), ("SS", None))
        self.assertEqual(clean_seat("", LARGE)[0], "SS")  # 申請時は SS で受付
        self.assertEqual(clean_seat("", SMALL)[0], "S")  # small は S のみ
        self.assertIsNotNone(clean_seat("SS", SMALL)[1])
        self.assertEqual(parse_show("1/5(月) 13時半", date(2026, 11, 1))[:2], (date(2027, 1, 5), "13:30"))
        self.assertEqual(classify_sheet("2026_1105_大公演_星の詩_申請").category, APPLICATION)


def person(kanji, kana, mail, phone, zip_, pref="大阪府", addr=None):
    return [kanji, kana, "", "", mail, phone, zip_, pref, addr or f"架空市{phone}", ""]


class PipelineTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        d = Path(self.tmp.name)
        a = person("架空 花子", "かくう はなこ", "Hanako@Example.com", 9000000001, "1000001")
        b = person("試験 太郎", "シケン タロウ", "taro@example.com", "080-0000-0002", "200-0002")
        a2 = person("架空　花子", "カクウ ハナコ", "", "090-0000-0001", "")  # 全角空白・メール無し
        c = person("試験 太郎", "シケン タロウ", "other@example.com", "070-0000-0003", "300-0003")  # 同姓同名別人

        with pd.ExcelWriter(d / "tickets.xlsx") as w:
            # タイトル行付き・列順バラバラのシート
            pd.DataFrame([["渤海 申込一覧", *[""] * 11], STD + ["Ticket", "Support"],
                          a + ["¥24,000", "5,000"], b + [12000, ""], ["合計", *[""] * 9, 36000, 5000]]
                         ).to_excel(w, sheet_name="2026_0121_大公演_渤海", header=False, index=False)
            pd.DataFrame([["名前", "フリガナ", "電話番号", "メール", "郵便番号", "チケット", "応援"],
                          [a2[0], a2[1], a2[5], "", "", 8000, 0]]
                         ).to_excel(w, sheet_name="2026_0415_小公演_春の夢", header=False, index=False)
            pd.DataFrame([STD + ["Ticket", "Support"], a + [30000, 10000], c + [6000, 0]]
                         ).to_excel(w, sheet_name="2026_0801_大公演_秋の詩", header=False, index=False)
        with pd.ExcelWriter(d / "events.xlsx") as w:
            pd.DataFrame([STD + ["Entry_fee"], a + [15000]]).to_excel(
                w, sheet_name="渤海_お茶会", header=False, index=False)
            pd.DataFrame([STD + ["total_cost"], b + [7000]]).to_excel(
                w, sheet_name="春の夢_親睦会", header=False, index=False)
            pd.DataFrame([STD + ["total_cost"], a + [3000]]).to_excel(
                w, sheet_name="2025_1220_クリスマス会", header=False, index=False)
            pd.DataFrame([STD + ["total_cost"], b + ["2,500"]]).to_excel(
                w, sheet_name="渤海_グッズ", header=False, index=False)
            pd.DataFrame([["メモ"]]).to_excel(w, sheet_name="Sheet1", header=False, index=False)
        cfg = d / "seat_prices.csv"
        cfg.write_text("category,seat_type,price,performance_sheet\nlarge,SS,12000,\nlarge,S,9000,\n",
                       encoding="utf-8")
        new = person("新規 由美", "シンキ ユミ", "yumi@example.com", "080-0000-0009", "400-0004")
        with pd.ExcelWriter(d / "apply.xlsx") as w:
            pd.DataFrame([STD + ["第1希望", "振替1", "振替2", "枚数", "席種", "サポート代（見込額）"],
                          a + ["11/12(木) 15:00", "11/13 11:00、11/14 15:00", "", 2, "", 20000],
                          new + ["11/15 13:00", "", "", "1枚", "S", ""]]
                         ).to_excel(w, sheet_name="2026_1105_大公演_星の詩_申請", header=False, index=False)
        self.prices = cfg
        self.result = run(sorted(d.glob("*.xlsx")), prices_path=d / "seat_prices.csv")

    def tearDown(self):
        self.tmp.cleanup()

    def test_events_and_links(self):
        ev = self.result.events.set_index("sheet_name")
        self.assertEqual(len(ev), 8)
        self.assertEqual(ev.at["渤海_お茶会", "linked_event_id"], ev.at["2026_0121_大公演_渤海", "event_id"])
        self.assertEqual(ev.at["春の夢_親睦会", "linked_event_id"], ev.at["2026_0415_小公演_春の夢", "event_id"])
        self.assertEqual(ev.at["渤海_グッズ", "linked_event_id"], ev.at["2026_0121_大公演_渤海", "event_id"])
        self.assertEqual(ev.at["2026_0801_大公演_秋の詩", "seq"], 2)
        self.assertTrue((self.result.review["sheet"] == "Sheet1").any())

    def test_identity(self):
        p = self.result.persons
        hanako = p[p["name_kana"] == "カクウ ハナコ"]
        self.assertEqual(len(hanako), 1, "表記ゆれ・連絡先の欠けがあっても同一人物にまとめる")
        self.assertEqual(hanako.iloc[0]["event_count"], 5, "申請は実績に数えない")
        self.assertEqual(hanako.iloc[0]["email"], "hanako@example.com")
        self.assertEqual(len(p[p["name_kanji"] == "試験 太郎"]), 2, "同姓同名でも連絡先が違えば別人")
        mr = self.result.match_review
        self.assertEqual(len(mr), 1)
        self.assertEqual(mr.iloc[0]["reason"], "名前が同じで連絡先が違う")

    def test_amounts(self):
        s = self.result.summary.set_index("customer_id")
        hanako_id = self.result.persons.query("name_kana == 'カクウ ハナコ'")["customer_id"].iloc[0]
        self.assertEqual(s.at[hanako_id, "large_ticket"], 24000 + 30000)
        self.assertEqual(s.at[hanako_id, "large_support"], 5000 + 10000)
        self.assertEqual(s.at[hanako_id, "small_ticket"], 8000)
        self.assertEqual(s.at[hanako_id, "support_total"], 15000)
        self.assertEqual(s.at[hanako_id, TEA_PARTY], 15000)
        self.assertEqual(s.at[hanako_id, ANNUAL], 3000)
        self.assertEqual(int(self.result.spending["amount"].sum()), 24000 + 5000 + 12000 + 8000 + 30000
                         + 10000 + 6000 + 15000 + 7000 + 3000 + 2500)

    def test_applications(self):
        req = self.result.application_requests
        pref = self.result.application_preferences
        p = self.result.persons.set_index("name_kana")
        hanako = req[req["customer_id"] == p.at["カクウ ハナコ", "customer_id"]].iloc[0]
        self.assertEqual((hanako["quantity"], hanako["seat_type_requested"], hanako["support_expected"]), (2, "SS", 20000))
        self.assertEqual((hanako["ticket_price"], hanako["ticket_amount"]), (12000, 24000))
        choices = pref[pref["request_id"] == hanako["request_id"]]
        self.assertEqual(list(zip(choices["show_date"], choices["show_time"])), [
            (date(2026, 11, 12), "15:00"), (date(2026, 11, 13), "11:00"), (date(2026, 11, 14), "15:00")])
        self.assertTrue(bool(p.at["シンキ ユミ", "is_new"]))
        # 見込額は実績の spending に入れない
        self.assertFalse((self.result.spending["category"] == APPLICATION).any())

    def test_select_events(self):
        sel = select_events(self.result.events, n_large=1, n_small=1, n_annual=1)
        self.assertEqual(set(sel["sheet_name"]), {
            "2026_0801_大公演_秋の詩", "2026_0415_小公演_春の夢", "春の夢_親睦会", "2025_1220_クリスマス会"})

    def test_write(self):
        with tempfile.TemporaryDirectory() as out:
            files = write_outputs(self.result, Path(out))
            self.assertTrue(all(f.exists() for f in files))


if __name__ == "__main__":
    unittest.main()


class SeatPriceEstimateTest(unittest.TestCase):
    def test_estimate_from_history_and_csv_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            a = person("架空 花子", "カクウ ハナコ", "hanako@example.com", "090-0000-0001", "100-0001")
            b = person("試験 太郎", "シケン タロウ", "taro@example.com", "080-0000-0002", "200-0002")
            cols = STD + ["席種", "枚数", "Ticket", "Support"]
            pd.DataFrame([a + ["SS", 2, 20000, 0], b + ["S", 1, 8000, 0]], columns=cols).to_csv(
                d / "2025_0601_大公演_旧作.csv", index=False)
            # 価格改定後: 新しい公演の単価を採用する
            pd.DataFrame([a + ["SS", 1, 12500, 3000], b + ["SS", 3, 37500, 0]], columns=cols).to_csv(
                d / "2026_0121_大公演_渤海.csv", index=False)
            # LINE からの申請 CSV（長い列名・cp932）
            pd.DataFrame([a + ["11/12 15:00", "2", "", "10000"]],
                         columns=STD + ["第1希望の公演日時", "希望枚数を入力してください", "席種", "サポート代（見込額）"]
                         ).to_csv(d / "2026_1105_大公演_星の詩_申請.csv", index=False, encoding="cp932")
            r = run(sorted(d.glob("*.csv")))
        est = r.seat_price_estimates.set_index("seat_type")
        self.assertEqual(est.at["SS", "price"], 12500)
        self.assertEqual(est.at["SS", "price_variants"], "10000 / 12500")
        self.assertEqual(est.at["S", "price"], 8000)
        req = r.application_requests.iloc[0]
        self.assertEqual((req["quantity"], req["seat_type_requested"], req["ticket_amount"], req["support_expected"]),
                         (2, "SS", 25000, 10000))
        self.assertTrue(r.review["message"].str.contains("推定値").any())


class MaskTest(unittest.TestCase):
    def test_mask_keeps_format_but_not_values(self):
        from takarazuka_data.mask import Masker, mask_sheet
        m = Masker()
        raw = pd.DataFrame([
            ["渤海 申込一覧", "", "", "", "", "", "", ""],
            ["name(漢字)", "name katakana", "mail", "phone-number", "post-number", "address 1", "Ticket", "備考"],
            ["架空　花子", "かくう はなこ", "Hanako@Example.com", 9000000001, "〒100-0001", "大阪府架空市1-2-3", 12000, "友人と"],
            ["架空花子", "カクウハナコ", "hanako@example.com", "０９０－００００－０００１", "", "", 8000, ""],
        ])
        out = mask_sheet(raw, m)
        flat = " ".join(map(str, out.values.ravel()))
        for secret in ("花子", "ハナコ", "はなこ", "Hanako", "0001", "架空市", "友人", "渤海"):
            self.assertNotIn(secret, flat)
        self.assertEqual(list(out.iloc[1]), list(raw.iloc[1]), "列名は残す")
        self.assertEqual(out.iat[2, 0][2], "　", "全角空白の位置を保つ")
        self.assertEqual(out.iat[2, 0].replace("　", ""), out.iat[3, 0], "同じ人は同じ架空名")
        self.assertTrue(all("ぁ" <= c <= "ゖ" or c == " " for c in out.iat[2, 1]), "ひらがなはひらがなのまま")
        self.assertIsInstance(out.iat[2, 3], int)  # 先頭 0 落ちの数値のまま
        self.assertEqual(len(str(out.iat[2, 3])), 10)
        self.assertRegex(out.iat[3, 3], r"^０[０-９]{2}－[０-９]{4}－[０-９]{4}$")
        self.assertRegex(out.iat[2, 4], r"^〒1\d{2}-\d{4}$")
        self.assertTrue(out.iat[2, 5].startswith("大阪府"))
        self.assertTrue(out.iat[2, 2].endswith("@example.com"))
        self.assertEqual(out.iat[2, 6], 12000, "金額は残す")
        self.assertNotEqual(Masker().name("架空 花子"), m.name("架空 花子"), "実行ごとに鍵が変わる")


class MasterTest(unittest.TestCase):
    """顧客マスター: 番号の維持、人の判定（同一/別人）の反映。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        self.raw, self.master, self.out = self.d / "raw", self.d / "master", self.d / "out"
        self.raw.mkdir()

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, sheet, rows, extra=("Ticket", "Support")):
        pd.DataFrame(rows, columns=STD + list(extra)).to_csv(self.raw / f"{sheet}.csv", index=False)

    def run_cli(self):
        from takarazuka_data.__main__ import main
        import contextlib, io
        with contextlib.redirect_stdout(io.StringIO()):
            main([str(self.raw), "-o", str(self.out), "--master", str(self.master),
                  "--links", str(self.d / "none.csv"), "--prices", str(self.d / "none.csv")])
        return (pd.read_csv(self.out / "persons.csv", dtype=str).fillna(""),
                pd.read_csv(self.out / "match_review.csv", dtype=str).fillna(""))

    def test_ids_are_stable_and_matching_uses_address_variants(self):
        a = person("髙橋 花子", "タカハシ ハナコ", "", "", "", addr="架空市中央1丁目2番3号")
        b = person("試験 太郎", "シケン タロウ", "taro@example.com", "", "")
        self.write("2026_0121_大公演_渤海", [b + [10000, 0], a + [12000, 1000]])
        persons, _ = self.run_cli()
        ids = dict(zip(persons["name_kana"], persons["customer_id"]))
        self.assertEqual(ids, {"シケン タロウ": "C00001", "タカハシ ハナコ": "C00002"})

        # 次回: 新しい人が先頭の行に来ても、既存の番号は変わらない。住所の表記ゆれ・異体字も同一人物
        new = person("新規 由美", "シンキ ユミ", "yumi@example.com", "", "")
        a2 = person("高橋 花子", "タカハシ ハナコ", "", "", "", addr="架空市中央1-2-3")
        self.write("2026_0415_小公演_春の夢", [new + [8000, 0], a2 + [8000, 500]])
        persons, _ = self.run_cli()
        ids = dict(zip(persons["name_kana"], persons["customer_id"]))
        self.assertEqual(ids, {"シケン タロウ": "C00001", "タカハシ ハナコ": "C00002", "シンキ ユミ": "C00003"})
        self.assertTrue(any((self.master / "backup").iterdir()), "上書き前にバックアップを取る")

    def test_review_decisions_are_applied_on_next_run(self):
        # 同じ電話で名前が違う（改姓など）→ 自動では結びつけず確認に回す
        a = person("架空 花子", "カクウ ハナコ", "", "090-0000-0001", "")
        a_new = person("試験 花子", "シケン ハナコ", "", "090-0000-0001", "")
        self.write("2026_0121_大公演_渤海", [a + [10000, 0]])
        self.write("2026_0415_小公演_春の夢", [a_new + [8000, 0]])
        persons, review = self.run_cli()
        self.assertEqual(len(persons), 2)
        self.assertIn("電話が同じで名前が違う", review.iloc[0]["reason"])

        review.loc[0, "decision"] = "同一"
        review.to_csv(self.out / "match_review.csv", index=False, encoding="cp932")  # Excel で保存した想定
        persons, review = self.run_cli()
        self.assertEqual(len(persons), 1, "「同一」と書いた組は統合される")
        self.assertEqual(persons.iloc[0]["customer_id"], "C00001")
        self.assertEqual(len(review), 0, "判定済みの組は再び聞かない")
        summary = pd.read_csv(self.out / "summary.csv")
        self.assertEqual(int(summary["total"].sum()), 18000)

    def test_line_id_links_people(self):
        a = person("架空 花子", "カクウ ハナコ", "", "", "")
        self.write("2026_0121_大公演_渤海", [a + ["U0123abcd", 10000, 0]], extra=("LINE ID", "Ticket", "Support"))
        a2 = person("架空花子", "", "", "", "")
        self.write("2026_0415_小公演_春の夢", [a2 + ["U0123abcd", 8000, 0]], extra=("LINE ID", "Ticket", "Support"))
        persons, _ = self.run_cli()
        self.assertEqual(len(persons), 1)
        self.assertEqual(persons.iloc[0]["line_id"], "U0123abcd")


class StandardizeTest(unittest.TestCase):
    """バラバラな既存シート → 指定形式 → 統合 までの流れ。"""

    def test_messy_sheets_to_standard_format(self):
        import contextlib, io
        from takarazuka_data.standardize import main as standardize
        from takarazuka_data.standardize import load_map
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            orig, raw, smap = d / "original", d / "raw", d / "sheet_map.csv"
            orig.mkdir()
            with pd.ExcelWriter(orig / "2026年 管理表.xlsx") as w:
                pd.DataFrame([
                    ["渤海 チケット申込者一覧", "", "", "", "", ""],
                    ["氏名", "フリガナ", "電話番号", "住所", "チケット代", "お花代"],
                    ["架空　花子", "かくう はなこ", 9000000001, "大阪府架空市1丁目2番3号", "¥24,000", "5,000円"],
                    ["合計", "", "", "", 24000, 5000],
                ]).to_excel(w, sheet_name="渤海 宝塚大劇場 2026.1.21", header=False, index=False)
                pd.DataFrame([["お名前", "カナ", "TEL", "参加費"],
                              ["架空花子", "カクウハナコ", "090-0000-0001", 15000]]
                             ).to_excel(w, sheet_name="渤海お茶会参加者", header=False, index=False)
                pd.DataFrame([["名前", "電話", "合計金額"], ["架空 花子", "09000000001", 3000]]
                             ).to_excel(w, sheet_name="2025年12月クリスマス会", header=False, index=False)
                pd.DataFrame([["メモ"]]).to_excel(w, sheet_name="メモ", header=False, index=False)
            args = [str(orig), "-o", str(raw), "--map", str(smap), "--aliases", str(d / "none.csv")]
            with contextlib.redirect_stdout(io.StringIO()):
                todo = standardize(args)
            m = load_map(smap).set_index("source_sheet")
            self.assertEqual(m.at["渤海 宝塚大劇場 2026.1.21", "target_sheet"], "2026_0121_大公演_渤海")
            self.assertEqual(m.at["渤海お茶会参加者", "target_sheet"], "渤海_お茶会")
            self.assertEqual(m.at["2025年12月クリスマス会", "target_sheet"], "2025_1201_クリスマス会")
            self.assertIn("日が不明", m.at["2025年12月クリスマス会", "note"])
            self.assertEqual(todo, 1, "分類できない「メモ」だけ記入待ち")

            out = pd.read_csv(raw / "2026_0121_大公演_渤海.csv", dtype=str).fillna("")
            self.assertEqual(list(out.columns), STD + ["Ticket", "Support"], "指定の列名・列順")
            self.assertEqual(len(out), 1, "タイトル行・合計行は除く")
            self.assertEqual(out.iloc[0]["phone-number"], "09000000001")
            self.assertEqual(out.iloc[0]["prefecture"], "大阪府")
            self.assertEqual(out.iloc[0]["name katakana"], "カクウ ハナコ")
            self.assertEqual((out.iloc[0]["Ticket"], out.iloc[0]["Support"]), ("24000", "5000"))

            # 人が sheet_map を直して再実行すると、修正が優先される
            fixed = load_map(smap)
            fixed.loc[fixed["source_sheet"] == "メモ", "category"] = "skip"
            fixed.loc[fixed["source_sheet"] == "2025年12月クリスマス会", "date"] = "20251220"
            fixed.to_csv(smap, index=False, encoding="cp932")
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(standardize(args), 0)
            self.assertTrue((raw / "2025_1220_クリスマス会.csv").exists())

            r = run(sorted(p for p in raw.glob("*.csv") if "1201" not in p.name))
        self.assertEqual(len(r.persons), 1, "整形後のシート同士で同一人物にまとまる")
        self.assertEqual(int(r.summary["total"].iloc[0]), 24000 + 5000 + 15000 + 3000)


class RosterSplitTest(unittest.TestCase):
    """名簿（連絡先）と金額のシートが分かれている場合。"""

    def test_amount_rows_without_contacts_link_to_roster(self):
        import contextlib, io
        from takarazuka_data.standardize import main as standardize
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            orig, raw = d / "original", d / "raw"
            orig.mkdir()
            with pd.ExcelWriter(orig / "母の管理表.xlsx") as w:
                pd.DataFrame([["氏名", "フリガナ", "電話番号", "住所"],
                              ["架空 花子", "カクウ ハナコ", "090-0000-0001", "大阪府架空市1-2-3"],
                              ["試験 太郎", "シケン タロウ", "080-0000-0002", "東京都架空区4-5-6"],
                              ["試験 太郎", "シケン タロウ", "070-0000-0003", "京都府架空市7-8-9"]]  # 同姓同名
                             ).to_excel(w, sheet_name="会員名簿", header=False, index=False)
                pd.DataFrame([["名前", "チケット代", "お花代"],
                              ["架空花子", 24000, 5000],      # 空白なし → 名簿の1人に結びつく
                              ["試験 太郎", 12000, 0],         # 同名2人 → 自動では結びつけない
                              ["名簿 にいない", 9000, 0]]
                             ).to_excel(w, sheet_name="渤海 宝塚大劇場 2026.1.21", header=False, index=False)
            with contextlib.redirect_stdout(io.StringIO()):
                todo = standardize([str(orig), "-o", str(raw), "--map", str(d / "map.csv"),
                                    "--aliases", str(d / "none.csv")])
            self.assertEqual(todo, 0)
            self.assertTrue((raw / "会員_名簿.csv").exists())
            r = run(sorted(raw.glob("*.csv")))

        p = r.persons.set_index("name_kana")
        hanako = p.loc["カクウ ハナコ"]
        self.assertEqual(hanako["phone"], "09000000001", "名簿の連絡先が付く")
        self.assertEqual(hanako["event_count"], 1, "名簿は参加イベントに数えない")
        s = r.summary.set_index("customer_id")
        self.assertEqual(s.at[hanako["customer_id"], "large_support"], 5000)
        self.assertEqual(len(r.persons), 5, "花子1 + 太郎(名簿2 + 金額のみ1) + 名簿にいない1")
        self.assertTrue(r.match_review["reason"].str.contains("連絡先のない方").any())
        self.assertTrue(r.review["message"].str.contains("同じ名前の人が複数").any())
        self.assertEqual(len(r.spending[r.spending["category"] == ROSTER]), 0)


class MaskLeakRegressionTest(unittest.TestCase):
    """実データで起きた漏れの再発防止（すべて架空の値）。"""

    SECRETS = ("花子", "ハナコ", "hanako", "太郎", "タロウ", "taro", "1234", "5678", "架空町", "夢見ヶ丘", "ゆめみ")

    def assertNoSecrets(self, out):
        flat = " ".join(map(str, out.values.ravel()))
        for s in self.SECRETS:
            self.assertNotIn(s, flat)
        for v in out.values.ravel():
            if isinstance(v, str) and "@" in v:
                self.assertTrue(v.endswith("@example.com"), "実在メールらしき値が残っている")

    def test_data_row_mistaken_for_header(self):
        # 列名の行がなく、LINE の注文明細のような長い文の行がヘッダと誤認されていた
        from takarazuka_data.mask import Masker, mask_sheet
        df = pd.DataFrame([
            ["2026-09-10 18:21:25", "12月27日11時00分", "・S席(10500円) 1枚・サポート代(任意) 1000円", 11500,
             "架空 花子", "カクウ ハナコ", "hanako.1234@gmail.com", "090-1234-5678", "665-0835", "兵庫県", "架空町1-2-3"],
            ["2026-09-11 10:00:00", "12月28日15時00分", "・S席(10500円) 1枚", 10500,
             "試験 太郎", "シケン タロウ", "taro5678@gmail.com", "080-5678-1234", "100-0001", "東京都", "夢見ヶ丘4-5"],
        ])
        out = mask_sheet(df, Masker())
        self.assertNoSecrets(out)
        self.assertEqual(out.iat[0, 0], "2026-09-10 18:21:25", "日時は残す")
        self.assertEqual(out.iat[0, 3], 11500, "金額は残す")

    def test_columns_shift_midway_and_names_next_to_headers(self):
        from takarazuka_data.mask import Masker, mask_sheet
        df = pd.DataFrame([
            ["お名前", "合計金額", "フリガナ", "メールアドレス", "電話番号", "郵便番号", "都道府県", "備考", "架空 花子"],
            ["架空 花子", 5000, "カクウ ハナコ", "hanako@gmail.com", "090-1234-5678", "665-0835", "兵庫県", "", ""],
            # ここから列がずれている（メール→電話列、氏名→郵便列、住所→都道府県列、電話が数値で備考列）
            ["試験 太郎", "シケン タロウ", "", "", "taro@gmail.com", "試験 太郎", "夢見ヶ丘4-5", 8012345678, ""],
            ["試験 太郎", "夢見ヶ丘4-5", "ゆめみ", "", "", "", "", "", ""],
        ])
        out = mask_sheet(df, Masker())
        self.assertNoSecrets(out)
        self.assertEqual(list(out.iloc[0, :7]), list(df.iloc[0, :7]), "本物の列名は残す")
        self.assertEqual(out.iat[1, 1], 5000)
        self.assertEqual(out.iat[1, 6], "兵庫県", "都道府県だと確かめられる値だけ残す")


class HeaderlessSheetTest(unittest.TestCase):
    """列名の行がないシート: 列の推測と、sheet_map の columns 列での指定。"""

    def make(self, d):
        import random
        random.seed(1)
        rows = []
        for i in range(1, 31):
            rows.append([i, None, f"架空 花{chr(0x5B50 + i)}", f"カクウ ハナ{chr(0x30A2 + i)}", f"0{random.randint(7, 9)}0-0000-{i:04d}",
                         f"{100 + i}-000{i % 10}", "大阪府", f"架空市{i}-2-3", 46028 + i % 3, 10000 + 500 * (i % 3),
                         1000 * (i % 4), "メモ" if i % 7 == 0 else None])
        rows.insert(15, ["No", None, "お名前", "フリガナ", "電話番号", "郵便番号", "都道府県", "住所", "日付", "代金", "サポート", "備考"])
        orig = d / "original"; orig.mkdir()
        with pd.ExcelWriter(orig / "2026管理.xlsx") as w:
            pd.DataFrame(rows).to_excel(w, sheet_name="121渤海チケット", header=False, index=False)
        return orig

    def run_std(self, d, orig):
        import contextlib, io
        from takarazuka_data.standardize import main as standardize
        with contextlib.redirect_stdout(io.StringIO()):
            standardize([str(orig), "-o", str(d / "raw"), "--map", str(d / "map.csv"), "--aliases", str(d / "none.csv")])

    def test_guess_and_override_columns(self):
        from takarazuka_data.standardize import load_map
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            orig = self.make(d)
            # 途中に列名の行が1行あるだけでは、先頭からの判定でヘッダとみなされないことを確認するため category を先に決める
            self.run_std(d, orig)
            m = load_map(d / "map.csv")
            m.loc[0, ["category", "date", "name"]] = ["large", "20260121", "渤海"]
            m.loc[0, "columns"] = ""  # 推測をやり直させる
            m.to_csv(d / "map.csv", index=False, encoding="utf-8-sig")
            self.run_std(d, orig)
            m = load_map(d / "map.csv")
            spec = m.loc[0, "columns"]
            self.assertIn("C=名前", spec)
            self.assertIn("D=カナ", spec)
            self.assertIn("E=電話", spec)
            self.assertIn("F=郵便番号", spec)
            self.assertIn("G=都道府県", spec)
            self.assertIn("H=住所1", spec)
            self.assertIn("J=チケット", spec)
            self.assertIn("K=サポート", spec)
            self.assertIn("列の指定", m.loc[0, "note"])
            out = pd.read_csv(d / "raw" / "2026_0121_大公演_渤海.csv", dtype=str).fillna("")
            self.assertEqual(len(out), 30, "途中の列名の行は飛ばす")
            self.assertEqual(out.iloc[0]["name katakana"], "カクウ ハナィ")
            self.assertEqual(out.iloc[0]["post-number"], "101-0001")

            # 人が指定を直すと、それが優先され、次回も上書きされない
            m.loc[0, "columns"] = "A=番号 C=名前 D=カナ E=電話 J=サポート K=チケット"
            m.to_csv(d / "map.csv", index=False, encoding="utf-8-sig")
            self.run_std(d, orig)
            self.assertEqual(load_map(d / "map.csv").loc[0, "columns"], "A=番号 C=名前 D=カナ E=電話 J=サポート K=チケット")
            out = pd.read_csv(d / "raw" / "2026_0121_大公演_渤海.csv", dtype=str).fillna("")
            self.assertEqual((out.iloc[0]["Ticket"], out.iloc[0]["Support"]), ("1000", "10500"))
            self.assertEqual(out.iloc[0]["post-number"], "", "指定しなかった列は取り込まない")


class SheetMapRobustnessTest(unittest.TestCase):
    """実データで見つかった sheet_map・列名まわりの不具合の再発防止（すべて架空の値）。"""

    def run_std(self, d, orig, extra=None):
        import contextlib, io
        from takarazuka_data.standardize import main as standardize
        with contextlib.redirect_stdout(io.StringIO()):
            standardize([str(orig), "-o", str(d / "raw"), "--map", str(d / "map.csv"), "--aliases", str(d / "none.csv")])

    def test_category_and_date_variants(self):
        from takarazuka_data.standardize import normalize_category, normalize_date
        for v, want in (("Small", "small"), (" LARGE ", "large"), ("大公演", "large"), ("お茶会", "tea_party"),
                        ("Tea Party", "tea_party"), ("取り込まない", "skip"), ("グッズ", "goods")):
            self.assertEqual(normalize_category(v), want)
        for v in ("20260121", "2026/1/21", "2026-01-21 00:00:00", "20,260,121", "２０２６年１月２１日"):
            self.assertEqual(normalize_date(v), "20260121", v)

    def test_unnamed_amount_column_header_without_name_column_and_aliases(self):
        from takarazuka_data.standardize import load_map
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            orig = d / "original"; orig.mkdir()
            names = [f"架空 花{chr(0x5B50 + i)}" for i in range(8)]
            kanas = [f"カクウ ハナ{chr(0x30A2 + 2 * i)}" for i in range(8)]
            with pd.ExcelWriter(orig / "管理.xlsx") as w:
                # 金額の列名が商品名（517グッズ）、ご協力金額（企画）
                pd.DataFrame({"お名前": names, "517グッズ": [1500 * (i % 3 + 1) for i in range(8)],
                              "フリガナ": kanas, "電話番号": [f"090-0000-{i:04d}" for i in range(8)]}
                             ).to_excel(w, sheet_name="グッズA", index=False)
                pd.DataFrame({"お名前": names, "ご協力金額": [1000 * (i % 3 + 1) for i in range(8)],
                              "フリガナ": kanas}).to_excel(w, sheet_name="企画B", index=False)
                # 列名の行はあるが、名前の列に列名がない
                pd.DataFrame([[None, None, "代金", "日付"]] + [
                    [names[i], kanas[i], 5500 + 1000 * i, 46028] for i in range(8)]
                ).to_excel(w, sheet_name="チケットC", header=False, index=False)
            self.run_std(d, orig)
            m = load_map(d / "map.csv").set_index("source_sheet")
            m.loc["グッズA", ["category", "date", "name"]] = ["Goods", "", "渤海"]       # 大文字・日本語でも可
            m.loc["企画B", ["category", "date", "name"]] = ["年間企画", "2026/3/31", "誕生日"]
            m.loc["チケットC", ["category", "date", "name"]] = ["Small", "2026.1.21", "渤海"]
            m.reset_index().to_csv(d / "map.csv", index=False, encoding="utf-8-sig")
            self.run_std(d, orig)
            m = load_map(d / "map.csv").set_index("source_sheet")
            self.assertEqual(m.at["グッズA", "target_sheet"], "渤海_グッズ")
            self.assertEqual(m.at["企画B", "target_sheet"], "2026_0331_誕生日")
            self.assertEqual(m.at["チケットC", "target_sheet"], "2026_0121_小公演_渤海")
            self.assertTrue(m.at["チケットC", "columns"], "名前の列がないシートも列を推測する")
            g = pd.read_csv(d / "raw" / "渤海_グッズ.csv", dtype=str)
            self.assertEqual(g["total_cost"].tolist()[:3], ["1500", "3000", "4500"])
            k = pd.read_csv(d / "raw" / "2026_0331_誕生日.csv", dtype=str)
            self.assertEqual(k["total_cost"].tolist()[:3], ["1000", "2000", "3000"])

    def test_city_only_addresses_do_not_match_people(self):
        from takarazuka_data.identity import contact_values
        base = {"prefecture": "兵庫県", "email": "", "phone": "", "postal_code": ""}
        self.assertEqual(contact_values({**base, "address1": "宝塚市"}), [])
        self.assertEqual([t for t, _ in contact_values({**base, "address1": "宝塚市旭町2-3-15"})], ["addr"])


class WideSheetSpecTest(unittest.TestCase):
    """横に長いシート: 2文字の列記号（AB など）の指定と、金額合計の表示。"""

    def test_two_letter_columns_and_totals(self):
        import contextlib, io
        from takarazuka_data.standardize import load_map, main as standardize
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            orig = d / "original"; orig.mkdir()
            rows = []
            for i in range(10):
                r = [None] * 35
                r[0] = f"架空 花{chr(0x5B50 + i)}"; r[1] = 2; r[2] = 1000; r[3] = 500   # A 名前 B 枚数 C サポート D 1枚あたり
                r[27] = f"カクウ ハナ{chr(0x30A2 + 2 * i)}"                              # AB カナ
                r[22] = 2; r[33] = 1000; r[34] = 500                                      # 別ブロックの同じ内容
                rows.append(r)
            with pd.ExcelWriter(orig / "管理.xlsx") as w:
                pd.DataFrame(rows).to_excel(w, sheet_name="606呂布東京チケット", header=False, index=False)
            buf = io.StringIO()
            args = [str(orig), "-o", str(d / "raw"), "--map", str(d / "map.csv"), "--aliases", str(d / "none.csv")]
            with contextlib.redirect_stdout(buf):
                standardize(args)
            m = load_map(d / "map.csv")
            m.loc[0, ["category", "date", "name", "columns"]] = ["large", "20260606", "呂布東京", "A=名前 AB=カナ B=枚数 C=サポート"]
            m.to_csv(d / "map.csv", index=False, encoding="utf-8-sig")
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                standardize(args)
            out = pd.read_csv(d / "raw" / "2026_0606_大公演_呂布東京.csv", dtype=str)
            self.assertEqual(len(out), 10)
            self.assertEqual(out.iloc[0]["name katakana"], "カクウ ハナア")
            self.assertEqual((out.iloc[0]["Support"], out.iloc[0]["枚数"]), ("1000", "2"))
            self.assertIn("Support 10,000円", buf.getvalue())


class RedownloadedFileTest(unittest.TestCase):
    """ダウンロードし直してファイル名が変わっても、sheet_map（人の修正）を引き継ぐ。"""

    def test_file_rename_keeps_map(self):
        import contextlib, io, os
        from takarazuka_data.standardize import load_map, main as standardize
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            orig = d / "original"; orig.mkdir()
            df = pd.DataFrame({"お名前": ["架空 花子", "試験 太郎", "新規 由美"], "合計金額": [1000, 2000, 3000],
                               "フリガナ": ["カクウ ハナコ", "シケン タロウ", "シンキ ユミ"]})
            with pd.ExcelWriter(orig / "2026成績仮 (1).xlsx") as w:
                df.to_excel(w, sheet_name="331誕生日企画", index=False)
            args = [str(orig), "-o", str(d / "raw"), "--map", str(d / "map.csv"), "--aliases", str(d / "none.csv")]
            with contextlib.redirect_stdout(io.StringIO()):
                standardize(args)
            m = load_map(d / "map.csv")
            m.loc[0, ["category", "date", "name"]] = ["annual", "20260331", "誕生日企画"]
            m.to_csv(d / "map.csv", index=False, encoding="utf-8-sig")
            # 新しいファイルをダウンロード（名前が違う）して、古い方を外す
            os.rename(orig / "2026成績仮 (1).xlsx", orig / "2026成績仮.xlsx")
            with contextlib.redirect_stdout(io.StringIO()):
                standardize(args)
            m = load_map(d / "map.csv")
            self.assertEqual(len(m), 1, "同じシートとして扱い、行を増やさない")
            self.assertEqual(m.loc[0, "target_sheet"], "2026_0331_誕生日企画")
            self.assertTrue((d / "raw" / "2026_0331_誕生日企画.csv").exists())


class RedRowAndHubTest(unittest.TestCase):
    def test_red_rows_are_excluded_and_totals_are_recorded(self):
        import contextlib, io
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill
        from takarazuka_data.standardize import load_map, main as standardize
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            orig = d / "original"; orig.mkdir()
            wb = Workbook(); ws = wb.active; ws.title = "606呂布東京チケット"
            ws.append(["お名前", "フリガナ", "サポート"])
            for i in range(6):
                ws.append([f"架空 花{chr(0x5B50 + i)}", f"カクウ ハナ{chr(0x30A2 + 2 * i)}", 1000 * (i + 1)])
            red = PatternFill("solid", start_color="FFFF0000", end_color="FFFF0000")
            for c in range(1, 9):  # 4行目（データの3人目）を行全体で赤く塗る
                ws.cell(row=4, column=c).fill = red
            ws.cell(row=6, column=1).fill = red  # 1セルだけ赤い行は除外しない
            wb.save(orig / "管理.xlsx")
            args = [str(orig), "-o", str(d / "raw"), "--map", str(d / "map.csv"), "--aliases", str(d / "none.csv")]
            with contextlib.redirect_stdout(io.StringIO()):
                standardize(args)
            m = load_map(d / "map.csv")
            m.loc[0, ["category", "date", "name"]] = ["large", "20260606", "呂布東京"]
            m.to_csv(d / "map.csv", index=False, encoding="utf-8-sig")
            with contextlib.redirect_stdout(io.StringIO()):
                standardize(args)
            out = pd.read_csv(d / "raw" / "2026_0606_大公演_呂布東京.csv", dtype=str)
            self.assertEqual(len(out), 5)
            self.assertNotIn("3000", out["Support"].tolist())
            self.assertEqual(load_map(d / "map.csv").loc[0, "totals"], "Ticket 0円 / Support 18,000円 / 赤い行を除外 1行")
            with contextlib.redirect_stdout(io.StringIO()):
                standardize(args + ["--include-red"])
            self.assertEqual(len(pd.read_csv(d / "raw" / "2026_0606_大公演_呂布東京.csv")), 6)

    def test_values_shared_by_many_people_are_not_reviewed(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            rows = [person(f"架空 花{chr(0x5B50 + i)}", f"カクウ ハナ{chr(0x30A2 + 2 * i)}", "shared@example.com",
                           f"090-0000-{i:04d}", f"10{i}-0001", addr=f"架空市{i}-2-3") + [1000, 0] for i in range(8)]
            rows += [person("なし", "ナシ", f"x{i}@example.com", f"080-1111-{i:04d}", f"20{i}-0002",
                            addr=f"試験市{i}-4-5") + [1000, 0] for i in range(7)]
            p = d / "2026_0121_大公演_渤海.csv"
            pd.DataFrame(rows, columns=STD + ["Ticket", "Support"]).to_csv(p, index=False)
            r = run([p])
        self.assertEqual(len(r.persons), 15)
        self.assertEqual(len(r.match_review), 0, "多数に共通するメール・名前は、確認の対象にしない")
        self.assertTrue(r.review["message"].str.contains("人以上に共通する値").any())


class OrphanRowTest(unittest.TestCase):
    def test_renamed_sheet_leaves_flagged_orphan_row(self):
        import contextlib, io
        from takarazuka_data.standardize import load_map, main as standardize
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            orig = d / "original"; orig.mkdir()
            df = pd.DataFrame({"お名前": ["架空 花子", "試験 太郎", "新規 由美"], "合計金額": [1000, 2000, 3000],
                               "フリガナ": ["カクウ ハナコ", "シケン タロウ", "シンキ ユミ"]})
            with pd.ExcelWriter(orig / "管理.xlsx") as w:
                df.to_excel(w, sheet_name="331誕生日企画", index=False)
            args = [str(orig), "-o", str(d / "raw"), "--map", str(d / "map.csv"), "--aliases", str(d / "none.csv")]
            with contextlib.redirect_stdout(io.StringIO()):
                standardize(args)
            with pd.ExcelWriter(orig / "管理.xlsx") as w:  # シート名を変更
                df.to_excel(w, sheet_name="20260331_誕生日企画", index=False)
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                standardize(args)
            m = load_map(d / "map.csv").set_index("source_sheet")
            self.assertIn("見つかりません", m.at["331誕生日企画", "note"])
            self.assertIn("20260331_誕生日企画", m.index)
            self.assertIn("元のシートが見つからない行: 1 件", buf.getvalue())


class PostSaleAndRenameTest(unittest.TestCase):
    def test_sheet_name_variants(self):
        from takarazuka_data.classify import POST_SALE
        info = classify_sheet("2026_0125_雨にじむ渤海_親睦会")
        self.assertEqual((info.category, info.name, info.event_date), (SOCIAL, "雨にじむ渤海", date(2026, 1, 25)))
        info = classify_sheet("呂布_事後販売")
        self.assertEqual((info.category, info.name), (POST_SALE, "呂布"))
        info = classify_sheet("2026_0121_small_雨にじむ渤海")
        self.assertEqual((info.category, info.name), (SMALL, "雨にじむ渤海"))

    def test_post_sale_is_not_linked_and_has_its_own_column(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            a = person("架空 花子", "カクウ ハナコ", "hanako@example.com", "090-0000-0001", "100-0001")
            pd.DataFrame([a + [1000, 0]], columns=STD + ["Ticket", "Support"]).to_csv(
                d / "2026_0606_大公演_呂布東京.csv", index=False)
            pd.DataFrame([a + [3000]], columns=STD + ["total_cost"]).to_csv(d / "呂布_事後販売.csv", index=False)
            r = run(sorted(d.glob("*.csv")))
        self.assertEqual(int(r.summary["post_sale"].iloc[0]), 3000)
        self.assertFalse(r.review["message"].str.contains("紐づく").any(), "事後販売は公演に紐づけないので警告も出ない")
        self.assertEqual(r.events.set_index("name").at["呂布", "linked_event_id"], "")

    def test_renamed_sheet_inherits_columns_from_old_row(self):
        import contextlib, io
        from takarazuka_data.standardize import load_map, main as standardize
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            orig = d / "original"; orig.mkdir()
            rows = [[i, f"架空 花{chr(0x5B50 + i)}", f"カクウ ハナ{chr(0x30A2 + 2 * i)}", 1000 * (i % 3)] for i in range(1, 9)]
            args = [str(orig), "-o", str(d / "raw"), "--map", str(d / "map.csv"), "--aliases", str(d / "none.csv")]

            def book(sheet):
                with pd.ExcelWriter(orig / "管理.xlsx") as w:
                    pd.DataFrame(rows).to_excel(w, sheet_name=sheet, header=False, index=False)

            book("404呂布宝塚チケット")
            with contextlib.redirect_stdout(io.StringIO()):
                standardize(args)
            m = load_map(d / "map.csv")
            m.loc[0, ["category", "date", "name", "columns"]] = ["large", "20260404", "呂布宝塚", "B=名前 C=カナ D=サポート"]
            m.to_csv(d / "map.csv", index=False, encoding="utf-8-sig")
            book("2026_0404_large_呂布宝塚")  # 標準の名前に変更
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                standardize(args)
            m = load_map(d / "map.csv")
            self.assertEqual(len(m), 1, "古い行は新しい行に置き換わる")
            self.assertEqual(m.loc[0, "columns"], "B=名前 C=カナ D=サポート")
            self.assertIn("引き継ぎました: 1 件", buf.getvalue())
            out = pd.read_csv(d / "raw" / "2026_0404_大公演_呂布宝塚.csv", dtype=str)
            self.assertEqual(len(out), 8)


class SplitSheetsTest(unittest.TestCase):
    """同じイベントが複数のシートに分かれている場合。"""

    def test_numbered_suffix_is_same_event(self):
        for name in ("雨にじむ渤海_親睦会2", "雨にじむ渤海_親睦会 (2)", "雨にじむ渤海_親睦会_3"):
            info = classify_sheet(name)
            self.assertEqual((info.category, info.name), (SOCIAL, "雨にじむ渤海"), name)
        info = classify_sheet("2026_0125_雨にじむ渤海_親睦会 (2)")
        self.assertEqual((info.category, info.name, info.event_date), (SOCIAL, "雨にじむ渤海", date(2026, 1, 25)))

    def test_split_sheets_are_merged_and_copies_are_flagged(self):
        import contextlib, io
        from takarazuka_data.standardize import load_map, main as standardize
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            orig = d / "original"; orig.mkdir()

            def people(start, n):
                return pd.DataFrame({"お名前": [f"架空 花{chr(0x5B50 + i)}" for i in range(start, start + n)],
                                     "フリガナ": [f"カクウ ハナ{chr(0x30A2 + 2 * i)}" for i in range(start, start + n)],
                                     "合計金額": [5500] * n})
            with pd.ExcelWriter(orig / "管理.xlsx") as w:
                people(0, 6).to_excel(w, sheet_name="雨にじむ渤海_親睦会", index=False)
                people(6, 4).to_excel(w, sheet_name="雨にじむ渤海_親睦会2", index=False)      # 別の人（追加分）
                people(0, 6).to_excel(w, sheet_name="雨にじむ渤海_親睦会3", index=False)      # 1枚目のコピー
            args = [str(orig), "-o", str(d / "raw"), "--map", str(d / "map.csv"), "--aliases", str(d / "none.csv")]
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                standardize(args)
            out = pd.read_csv(d / "raw" / "雨にじむ渤海_親睦会.csv", dtype=str)
            self.assertEqual(len(out), 16, "3枚を1つのイベントにまとめる")
            self.assertIn("3枚のシートを1つにまとめました", buf.getvalue())
            self.assertIn("重なっているシートの組: 1 組", buf.getvalue())
            m = load_map(d / "map.csv").set_index("source_sheet")
            self.assertIn("重なっている", m.at["雨にじむ渤海_親睦会3", "note"])
            self.assertIn("重なっている", m.at["雨にじむ渤海_親睦会", "note"])
            self.assertNotIn("重なっている", m.at["雨にじむ渤海_親睦会2", "note"])
