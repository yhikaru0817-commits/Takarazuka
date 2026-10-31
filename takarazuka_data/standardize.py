"""既存のバラバラなシートを、決められた形式（シート名・列順）に整形する。ローカル専用。

流れ:
    1. data/original/ の全シートについて、種別・日付・名前を推測して data/sheet_map.csv に書く
       （すでに書かれている行は人の修正を優先して残す）
    2. sheet_map.csv の内容に従って、各シートを標準の列に揃えて data/raw/<標準シート名>.csv に出力
       （確認用に data/standardized.xlsx にもまとめる）
    3. その後 python -m takarazuka_data で顧客マスター作成・統合を行う

使い方: python -m takarazuka_data.standardize [元フォルダ] [-o 出力フォルダ] [--map sheet_map.csv]
"""
from __future__ import annotations

import argparse
import re
import shutil
from datetime import datetime
import unicodedata
from pathlib import Path

import pandas as pd

from .classify import (
    ANNUAL, GOODS, LARGE, POST_SALE, REFUND, ROSTER, SMALL, SOCIAL, TEA_PARTY, UNKNOWN, classify_sheet, name_key,
    normalize_text,
)
from .columns import (
    AMOUNT_FIELDS, EXPECTED_AMOUNTS, PERSON_FIELDS, PREFECTURES, clean_amount, clean_prefecture, is_blank,
    load_aliases, map_header,
)
from .pipeline import _is_total_row, find_header, normalize_row, read_sheets

CATEGORIES = [LARGE, SMALL, TEA_PARTY, SOCIAL, ANNUAL, GOODS, POST_SALE, REFUND, ROSTER, "skip"]
CATEGORY_ALIASES = {
    "大公演": LARGE, "大劇場": LARGE, "小公演": SMALL, "別箱": SMALL, "お茶会": TEA_PARTY, "親睦会": SOCIAL,
    "グッズ": GOODS, "事後販売": POST_SALE, "返金": REFUND, "年間企画": ANNUAL, "企画": ANNUAL, "名簿": ROSTER,
    "取り込まない": "skip", "除外": "skip", "スキップ": "skip", "対象外": "skip",
}


def normalize_category(v) -> str:
    """'Small' や ' large ' や日本語（大公演・お茶会など）でも受け付ける。"""
    s = re.sub(r"[\s\-]+", "_", unicodedata.normalize("NFKC", str(v)).strip()).lower()
    return CATEGORY_ALIASES.get(s, s)


def normalize_date(v) -> str:
    """20260121 / 2026/1/21 / 2026-01-21 / 20,260,121 のどれでも 20260121 にする。"""
    s = unicodedata.normalize("NFKC", str(v)).strip()
    m = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", s)
    if m:
        return f"{int(m.group(1)):04d}{int(m.group(2)):02d}{int(m.group(3)):02d}"
    return re.sub(r"\D", "", s)
CATEGORY_LABEL = {LARGE: "大公演", SMALL: "小公演"}
SUFFIX_LABEL = {TEA_PARTY: "お茶会", SOCIAL: "親睦会", GOODS: "グッズ", ROSTER: "名簿", POST_SALE: "事後販売"}
MAP_COLUMNS = ["source_file", "source_sheet", "category", "date", "name", "columns", "target_sheet", "rows", "totals", "note"]

# 指定の列名と順番
OUTPUT_HEADERS = {
    "name_kanji": "name(漢字)", "name_kana": "name katakana", "family_name": "Family name",
    "first_name": "First name", "email": "mail", "phone": "phone-number", "postal_code": "post-number",
    "prefecture": "prefecture", "address1": "address 1", "address2": "address 2",
    "ticket": "Ticket", "support": "Support", "entry_fee": "Entry_fee", "total_cost": "total_cost",
}
OPTIONAL_HEADERS = {"line_id": "LINE ID", "seat_type": "席種", "quantity": "枚数"}  # 元にあれば残す
BASE_FIELDS = [f for f in PERSON_FIELDS if f != "line_id"]

# 種別の推測に使う語（シート名・ファイル名から）
_GUESS_WORDS = [
    (REFUND, ["返金"]),
    (TEA_PARTY, ["お茶会", "茶会"]),
    (SOCIAL, ["親睦会"]),
    (POST_SALE, ["事後販売"]),
    (GOODS, ["グッズ", "goods"]),
    (LARGE, ["大公演", "大劇場", "宝塚大劇場", "東京宝塚劇場", "本公演", "large"]),
    (SMALL, ["小公演", "別箱", "小劇場", "バウ", "梅田", "ドラマシティ", "全国ツアー", "博多座", "御園座",
             "KAAT", "赤坂", "東京建物", "small"]),
    (ROSTER, ["名簿", "住所録", "会員一覧", "会員リスト", "顧客リスト", "顧客一覧", "連絡先"]),  # 最後に判定
]
_DATE_PATTERNS = [
    re.compile(r"(20\d{2})[年/.\-_ ](\d{1,2})[月/.\-_ ](\d{1,2})日?(?!\d)"),
    re.compile(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)"),
    re.compile(r"(20\d{2})[年/.\-_ ](\d{1,2})月?()(?!\d)"),  # 日がない
]


def guess(sheet_name: str, file_name: str, fields: set) -> dict:
    """シート名・ファイル名・列から 種別 / 日付 / 名前 を推測する。"""
    info = classify_sheet(sheet_name)
    if info.category != UNKNOWN and not info.error:  # すでに命名規則どおり
        return {"category": info.category, "date": info.event_date.strftime("%Y%m%d") if info.event_date else "",
                "name": info.name, "note": ""}

    text = normalize_text(sheet_name)
    both = text + " " + normalize_text(Path(file_name).stem)
    notes = []
    category = ""
    for cat, words in _GUESS_WORDS:
        if any(w.lower() in both.lower() for w in words):
            category = cat
            break
    if not category:  # 列から推測
        if "entry_fee" in fields:
            category = TEA_PARTY
        elif "ticket" in fields or "support" in fields:
            category = LARGE
            notes.append("大公演/小公演を確認")
        elif "total_cost" in fields:
            category = ANNUAL
            notes.append("種別を確認（年間企画と推測）")
        elif fields & {"email", "phone", "postal_code", "address1"}:
            category = ROSTER
            notes.append("金額の列がないため名簿と推測")
        else:
            notes.append("種別を記入してください")
    else:
        notes.append("推測")

    date = ""
    for src in (text, both):
        for pat in _DATE_PATTERNS:
            m = pat.search(src)
            if m:
                y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3) or 1)
                if 1 <= mo <= 12 and 1 <= d <= 31:
                    date = f"{y:04d}{mo:02d}{d:02d}"
                    if not m.group(3):
                        notes.append("日が不明のため1日としました")
                    break
        if date:
            break
    if not date and category in (LARGE, SMALL, ANNUAL):
        notes.append("日付を記入してください")

    name = text
    for w in sorted((w for _, ws in _GUESS_WORDS for w in ws), key=len, reverse=True):  # 長い語から消す
        name = re.sub(re.escape(w), " ", name, flags=re.IGNORECASE)
    for pat in _DATE_PATTERNS:
        name = pat.sub(" ", name)
    name = re.sub(r"[_\s　()（）\[\]【】「」・\-]+", " ", name).strip()
    for w in ("申込み", "申込", "名簿", "一覧", "リスト", "参加者", "購入者", "チケット", "公演",
              "花組", "月組", "雪組", "星組", "宙組", "専科"):
        name = name.replace(w, " ")
    name = re.sub(r"\s+", " ", name).strip()
    if not name and category == ROSTER:
        name = "会員"
    if not name:
        notes.append("名前を記入してください")
    return {"category": category, "date": date, "name": name, "note": "・".join(notes)}


def target_sheet(category: str, date: str, name: str) -> str:
    """標準のシート名を作る。必要な項目が欠けていれば空文字。"""
    name = normalize_text(name).replace("_", " ")
    if category in (LARGE, SMALL):
        return f"{date[:4]}_{date[4:8]}_{CATEGORY_LABEL[category]}_{name}" if len(date) == 8 and name else ""
    if category == REFUND:  # 返金は、同じ公演名のグッズにマイナスの金額として加える
        return f"{name}_グッズ" if name else ""
    if category in SUFFIX_LABEL:
        return f"{name}_{SUFFIX_LABEL[category]}" if name else ""
    if category == ANNUAL:
        return f"{date[:4]}_{date[4:8]}_{name}" if len(date) == 8 and name else ""
    return ""


# ---------- 列名の行がないシート: 列の指定 ----------
# sheet_map.csv の columns 列に「A=番号 B=- C=名前 D=カナ ...」のように書く（A,B,C は Excel の列）。
# 「-」や番号・日付・人数・備考などは取り込まない。
IGNORE_LABELS = {"", "-", "ー", "番号", "日付", "日時", "人数", "備考", "メモ", "不明", "?", "？"}
_SPEC_TOKEN = re.compile(r"([A-Za-z]{1,2})\s*[=:：＝]\s*([^\s,、]*)")


def col_letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def col_index(letter: str) -> int:
    n = 0
    for ch in letter.upper():
        n = n * 26 + ord(ch) - 64
    return n - 1


def parse_columns(spec: str, ncols: int, issue) -> dict:
    """'A=- B=名前 C=カナ' → {列番号: 項目}"""
    mapping = {}
    for letter, label in _SPEC_TOKEN.findall(unicodedata.normalize("NFKC", spec)):
        i = col_index(letter)
        label = label.strip()
        if label in IGNORE_LABELS:
            continue
        field = map_header(label, contains=False) or {"住所": "address1", "建物": "address2", "金額": None}.get(label)
        if not field:
            issue(f"列の指定「{letter}={label}」が分かりません（名前・カナ・電話・メール・郵便番号・都道府県・住所1・住所2・"
                  "チケット・サポート・参加費・合計・枚数・席種 のどれかにしてください）")
            continue
        if i >= ncols:
            issue(f"列の指定「{letter}」はシートにない列です")
            continue
        mapping[i] = field
    return mapping


_KATA = re.compile(r"^[\u30A0-\u30FF\u3040-\u309F\uFF65-\uFF9F\s　・]+$")
_POSTAL = re.compile(r"^〒?\s*\d{3}[-－ー]?\d{4}$")
_PHONE = re.compile(r"^\+?[\d\-－()（）\s]{10,15}$")


def guess_columns(df: pd.DataFrame, category: str) -> str:
    """値の形から各列の種類を推測して「A=名前 B=カナ ...」を作る（推測なので人が確認する）。"""
    labels = []
    amount_cols = []
    used = set()
    for ci, c in enumerate(df.columns):
        vals = [v for v in df[c].tolist() if not is_blank(v)][:80]
        label = "-"
        if len(vals) >= max(3, len(df) * 0.2):
            txt = [unicodedata.normalize("NFKC", str(v)).strip() for v in vals]
            share = lambda f: sum(1 for t in txt if f(t)) / len(txt)
            nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
            num_share = len(nums) / len(vals)
            if share(lambda t: "@" in t) > .6:
                label = "メール"
            elif share(lambda t: bool(_POSTAL.match(t))) > .6 or (num_share > .8 and all(
                    100000 <= n <= 9999999 for n in nums) and "postal" not in used and "phone" in used):
                label = "郵便番号"
            elif share(lambda t: bool(_PHONE.match(t)) and re.sub(r"\D", "", t)[:1] in "0" or
                       (t.isdigit() and len(t) in (9, 10))) > .6:
                label = "電話"
            elif share(lambda t: clean_prefecture(t) in PREFECTURES) > .6:
                label = "都道府県"
            elif share(lambda t: bool(_KATA.match(t))) > .7 and "name_kana" not in used:
                label = "カナ"
            elif num_share > .8:
                ints = [int(n) for n in nums]
                if all(1 <= n <= 9 for n in ints):
                    label = "枚数" if category in (LARGE, SMALL) else "人数"
                elif all(30000 <= n <= 60000 for n in ints):
                    label = "日付"  # Excel の日付シリアル値
                elif ints == sorted(ints) and len(set(ints)) == len(ints) and ints[0] <= 2:
                    label = "番号"
                elif sum(1 for n in ints if n % 100 == 0) / len(ints) > .7:
                    label = "金額"
                    amount_cols.append(len(labels))
            elif share(lambda t: bool(re.search(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}", t))) > .6:
                label = "日付"
            elif share(lambda t: bool(re.search(r"\d", t)) and len(t) >= 4) > .6:
                label = "住所1" if "address1" not in used else "住所2"
            elif share(lambda t: 2 <= len(t) <= 12 and not re.search(r"\d", t)) > .7 and "name_kanji" not in used:
                label = "名前"
        field = map_header(label, contains=False) or {"住所1": "address1", "住所2": "address2"}.get(label)
        if field:
            used.add(field)
        labels.append(label)
    # 金額列: 公演は チケット→サポート、お茶会は参加費、その他は合計
    names = {LARGE: ["チケット", "サポート"], SMALL: ["チケット", "サポート"], TEA_PARTY: ["参加費"]}.get(category, ["合計"])
    for n, i in enumerate(amount_cols):
        labels[i] = names[n] if n < len(names) else "-"
    while labels and labels[-1] == "-":
        labels.pop()
    return " ".join(f"{col_letter(i)}={l}" for i, l in enumerate(labels))


# ---------- 赤く塗られた行（キャンセルなど）の除外 ----------
RED_MIN_CELLS = 5  # 1行のうち、赤く塗られたセルがこの数以上なら「赤い行」


def _is_red(rgb) -> bool:
    h = str(rgb)[-6:]
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except ValueError:
        return False
    return r >= 0xD0 and g <= 0x60 and b <= 0x60


def load_red_rows(path: Path) -> dict:
    """{シート名: 赤い行の集合（0始まり。pandas の行番号と一致）}。xlsx のみ。"""
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return {}
    from openpyxl import load_workbook
    out = {}
    wb = load_workbook(path)
    for ws in wb.worksheets:
        rows = set()
        for i, row in enumerate(ws.iter_rows()):
            red = 0
            for c in row:
                f = c.fill
                if f is not None and f.fill_type == "solid" and f.fgColor is not None \
                        and f.fgColor.type == "rgb" and _is_red(f.fgColor.rgb):
                    red += 1
            if red >= RED_MIN_CELLS:
                rows.add(i)
        out[ws.title] = rows
    return out


def _overlap(a: pd.DataFrame, b: pd.DataFrame) -> float:
    """2つのシートで「同じ人・同じ金額」の行がどれだけ重なっているか（0〜1）。コピーや整理版の見つけ出し用。"""
    cols = [c for c in ("Ticket", "Support", "Entry_fee", "total_cost") if c in a.columns and c in b.columns]
    if not cols:
        return 0.0

    def keys(df):
        out = set()
        for _, r in df.iterrows():
            who = name_key(r.get("name(漢字)") or r.get("name katakana") or "")
            if who:
                out.add((who, tuple(r[c] for c in cols)))
        return out

    ka, kb = keys(a), keys(b)
    return len(ka & kb) / min(len(ka), len(kb)) if ka and kb else 0.0


def file_key(name) -> str:
    """ファイル名の照合用キー。ダウンロードし直して「 (1)」が付いたり消えたりしても、同じファイルとみなす。"""
    stem = unicodedata.normalize("NFKC", Path(str(name)).stem)
    return re.sub(r"\s*\(\d+\)\s*$", "", stem).strip().lower()


def load_map(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=MAP_COLUMNS)
    for enc in ("utf-8-sig", "cp932"):
        try:
            return pd.read_csv(path, dtype=str, encoding=enc).fillna("").reindex(columns=MAP_COLUMNS, fill_value="")
        except UnicodeDecodeError:
            continue
    raise ValueError(f"文字コードを判定できません: {path.name}")


def standardize_sheet(df: pd.DataFrame, category: str, issue, spec: str = "", skip_rows=frozenset(),
                      stats=None) -> tuple[pd.DataFrame | None, list[str]]:
    """1シートを標準の列に揃える。spec（列の指定）があればそれを優先。戻り値: (整形後の表, 未対応の列名)"""
    if spec.strip():
        by_index = parse_columns(spec, len(df.columns), issue)
        mapping = {df.columns[i]: f for i, f in by_index.items()}
        header_row, unmapped = -1, []
        if len(set(mapping.values())) < 1:
            issue("列の指定から取り込める列がありません")
            return None, []
    else:
        header_row, mapping = find_header(df)
        if header_row is None:
            issue("列名の行がありません。sheet_map.csv の columns 列で、各列が何かを指定してください")
            return None, []
        unmapped = [str(h) for c, h in df.iloc[header_row].items() if c not in mapping and not is_blank(h)]
    amounts = EXPECTED_AMOUNTS[category]
    found = set(mapping.values())
    if header_row != -1 and len(amounts) == 1 and not (found & set(AMOUNT_FIELDS)):
        # 金額の列名が「517グッズ」のように独自の名前のとき: 数字だけの列が1つだけなら金額とみなす
        body = df.iloc[header_row + 1:]
        cand = []
        for c in df.columns:
            if c in mapping:
                continue
            vals = [v for v in body[c].tolist() if not is_blank(v)]
            nums = [v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool)]
            if len(vals) >= 3 and len(nums) / len(vals) >= 0.7:
                cand.append(c)
        if len(cand) > 1:  # 数字の列が複数あるときは、列名に「金・額・返・代・費・円」を含む列だけを候補にする
            cand = [c for c in cand if re.search(r"金|額|返|代|費|円", str(df.iat[header_row, df.columns.get_loc(c)]))]
        if len(cand) == 1:
            head = df.iat[header_row, df.columns.get_loc(cand[0])]
            mapping = {**mapping, cand[0]: amounts[0]}
            found = set(mapping.values())
            unmapped = [u for u in unmapped if u != str(head)]
            issue(f"列「{head}」を {OUTPUT_HEADERS[amounts[0]]} として扱いました（違う場合は columns で指定）")
    for f in amounts:
        if f not in found:
            # 金額列が別名なら、他の金額列で代用（例: お茶会シートに「合計」しかない）
            alt = next((a for a in AMOUNT_FIELDS if a in found and a not in amounts), None)
            if alt and len(amounts) == 1:
                mapping = {c: (f if v == alt else v) for c, v in mapping.items()}
                issue(f"金額列 {alt} を {OUTPUT_HEADERS[f]} として扱いました")
            else:
                issue(f"金額列 {OUTPUT_HEADERS[f]} が見つかりません")
    optional = [f for f in OPTIONAL_HEADERS if f in set(mapping.values())] if category in (LARGE, SMALL) else \
        (["line_id"] if "line_id" in set(mapping.values()) else [])

    out = []
    for i in range(header_row + 1, len(df)):
        if header_row == -1 and sum(1 for v in df.iloc[i] if map_header(v, contains=False)) >= 2:
            continue  # 列の指定を使うとき、途中にある列名の行は飛ばす
        raw = {f: df.iat[i, df.columns.get_loc(c)] for c, f in mapping.items()}
        if all(is_blank(raw.get(f)) for f in PERSON_FIELDS) or _is_total_row(raw):
            continue
        if i in skip_rows and not (is_blank(raw.get("name_kanji")) and is_blank(raw.get("name_kana"))):
            if stats is not None:
                stats["red"] = stats.get("red", 0) + 1
            continue  # 赤く塗られた行は除外
        rec = normalize_row(raw, lambda level, msg, r=i + 1: issue(f"{r}行目: {msg}"))
        if not rec["name_kanji"] and not rec["name_kana"]:
            issue(f"{i + 1}行目: 氏名がないため除外しました")
            continue
        row = {OUTPUT_HEADERS[f]: rec[f] for f in BASE_FIELDS}
        for f in amounts:
            row[OUTPUT_HEADERS[f]] = -abs(rec[f]) if category == REFUND else rec[f]
        for f in optional:
            v = raw.get(f)
            if f == "quantity":
                v = clean_amount(v)[0] if not is_blank(v) else ""
            elif f == "line_id":
                v = rec["line_id"]
            row[OPTIONAL_HEADERS[f]] = "" if is_blank(v) else normalize_text(v)
        out.append(row)
    cols = [OUTPUT_HEADERS[f] for f in BASE_FIELDS + amounts] + [OPTIONAL_HEADERS[f] for f in optional]
    return pd.DataFrame(out, columns=cols), unmapped


def main(argv=None):
    ap = argparse.ArgumentParser(description="既存のシートを決められた形式に整形します（ネット通信なし）")
    ap.add_argument("input", nargs="?", default="data/original", help="既存のシートを置いたフォルダ")
    ap.add_argument("-o", "--output", default="data/raw", help="整形後のCSVの出力先")
    ap.add_argument("--map", default="data/sheet_map.csv", help="シートの対応表")
    ap.add_argument("--aliases", default="config/column_aliases.csv", help="独自の列名の対応")
    ap.add_argument("--include-red", action="store_true", help="赤く塗られた行も取り込む（既定は除外）")
    args = ap.parse_args(argv)

    src, dst, map_path = Path(args.input), Path(args.output), Path(args.map)
    if not src.exists():
        src.mkdir(parents=True)
        raise SystemExit(f"{src} フォルダを作りました。既存のシート（Excel/CSV）をそのまま入れて、もう一度実行してください")
    n_alias = load_aliases(args.aliases)
    if n_alias:
        print(f"独自の列名を {n_alias} 件読み込みました")
    paths = sorted(p for p in src.iterdir()
                   if p.suffix.lower() in (".xlsx", ".xlsm", ".xls", ".csv") and not p.name.startswith("~$"))

    sheets = list(read_sheets(paths))
    smap = load_map(map_path)
    known = {(file_key(r["source_file"]), r["source_sheet"]) for _, r in smap.iterrows()}
    added = []
    for sheet_name, df, file_name in sheets:
        if (file_key(file_name), str(sheet_name)) in known:
            continue
        _, mapping = find_header(df)
        g = guess(str(sheet_name), file_name, set(mapping.values()))
        added.append({"source_file": file_name, "source_sheet": str(sheet_name), **g})
    inherited = 0
    if added:
        present = {(file_key(f), str(s)) for s, _, f in sheets}
        orphans_idx = [i for i, o in smap.iterrows()
                       if (file_key(o["source_file"]), o["source_sheet"]) not in present
                       and normalize_category(o["category"]) not in ("skip", "")]
        used = set()
        for a in added:
            cat_a = normalize_category(a.get("category", ""))
            if not cat_a:
                continue
            cands = []
            for i in orphans_idx:
                o = smap.loc[i]
                if i in used or normalize_category(o["category"]) != cat_a:
                    continue
                if name_key(o["name"]) != name_key(a.get("name", "")):
                    continue
                if cat_a in (LARGE, SMALL, ANNUAL) and normalize_date(o["date"]) != normalize_date(a.get("date", "")):
                    continue
                cands.append(i)
            if len(cands) == 1:  # 種別・名前（・日付）が同じ古い行が1つだけなら、同じシートの名前変更とみなす
                used.add(cands[0])
                if str(smap.loc[cands[0], "columns"]).strip():
                    a["columns"] = smap.loc[cands[0], "columns"]
                a["note"] = "・".join(x for x in (str(a.get("note", "")), "古いシート名の設定を引き継ぎました") if x)
                inherited += 1
        smap = smap.drop(index=list(used))
        smap = pd.concat([smap, pd.DataFrame(added)], ignore_index=True).reindex(columns=MAP_COLUMNS).fillna("")
    smap["category"] = smap["category"].map(normalize_category)
    smap["date"] = smap["date"].map(lambda d: normalize_date(d) if str(d).strip() else "")
    # 列名の行がないシートには、列の推測を columns 列に書く（空欄のときだけ。人の修正は上書きしない）
    data, newest, file_of = {}, {}, {}
    red_cache, total_red = {}, 0
    for s_name, s_df, f_name in sheets:  # 同じシートが新旧2つのファイルにあれば、新しい方（更新日時）を使う
        key = (file_key(f_name), str(s_name))
        mtime = (src / f_name).stat().st_mtime if (src / f_name).exists() else 0
        if key in data:
            print(f"注意: 同じシート「{s_name}」が複数のファイルにあります。新しい方を使います（古いファイルは data/original から外してください）")
        if key not in data or mtime >= newest[key]:
            data[key], newest[key] = s_df, mtime
            file_of[key] = f_name
    guessed_cols = 0
    for idx, r in smap.iterrows():
        df = data.get((file_key(r["source_file"]), r["source_sheet"]))
        if df is None or r["columns"].strip() or r["category"].strip() == "skip":
            continue
        hr, mp = find_header(df)
        if hr is None or not ({"name_kanji", "name_kana", "family_name"} & set(mp.values())):
            smap.at[idx, "columns"] = guess_columns(df, r["category"].strip())
            smap.at[idx, "note"] = "・".join(x for x in (r["note"], "列の指定（columns）を確認") if x)
            guessed_cols += 1
    smap["target_sheet"] = [target_sheet(c.strip(), re.sub(r"\D", "", d), n.strip())
                            for c, d, n in zip(smap["category"], smap["date"], smap["name"])]

    # 整形
    report, merged, orphans, merged_idx = [], {}, 0, {}
    for idx, r in smap.iterrows():
        def issue(msg, r=r):
            report.append({"source_file": r["source_file"], "source_sheet": r["source_sheet"], "message": msg})
        cat = r["category"].strip()
        if cat == "skip":
            continue
        if (file_key(r["source_file"]), r["source_sheet"]) not in data:
            issue("元のシートが見つかりません（sheet_map.csv の行を消して構いません）")
            if "見つかりません" not in r["note"]:
                smap.at[idx, "note"] = "・".join(x for x in (r["note"], "元のシートが見つかりません（この行は削除してよい）") if x)
            orphans += 1
            continue
        if cat not in CATEGORIES or not r["target_sheet"]:
            issue("sheet_map.csv の category / date / name を記入してください（取り込まない場合は category に skip）")
            continue
        key = (file_key(r["source_file"]), r["source_sheet"])
        skip_rows = frozenset()
        if not args.include_red:
            path = src / file_of[key]
            if path not in red_cache:
                try:
                    red_cache[path] = load_red_rows(path)
                except Exception:  # 色が読めなくても整形は続ける
                    red_cache[path] = {}
                    issue("セルの色を読めなかったため、赤い行の除外はできませんでした")
            skip_rows = frozenset(red_cache[path].get(r["source_sheet"], ()))
        stats = {}
        out, unmapped = standardize_sheet(data[key], cat, issue, r["columns"], skip_rows, stats)
        if unmapped:
            issue("使わなかった列: " + ", ".join(unmapped) + "（必要なら column_aliases.csv に追加）")
        if out is None:
            continue
        smap.at[idx, "rows"] = str(len(out))
        tot = [f"{col} {int(pd.to_numeric(out[col], errors='coerce').fillna(0).sum()):,}円"
               for col in ("Ticket", "Support", "Entry_fee", "total_cost") if col in out.columns]
        if stats.get("red"):
            tot.append(f"赤い行を除外 {stats['red']}行")
            total_red += stats["red"]
        smap.at[idx, "totals"] = " / ".join(tot)
        merged.setdefault(r["target_sheet"], []).append(out)
        merged_idx.setdefault(r["target_sheet"], []).append(idx)

    # 複数のシートを1つのイベントにまとめた場合: まとめた枚数を表示し、同じ内容が重なっていれば警告する
    dup_pairs, combined = 0, []
    for target, parts in merged.items():
        if len(parts) < 2:
            continue
        combined.append((target, len(parts)))
        idxs = merged_idx[target]
        for i in range(len(parts)):
            for j in range(i + 1, len(parts)):
                if _overlap(parts[i], parts[j]) >= 0.8:
                    dup_pairs += 1
                    for k in (idxs[i], idxs[j]):
                        if "重なっている" not in smap.at[k, "note"]:
                            smap.at[k, "note"] = "・".join(x for x in (
                                smap.at[k, "note"], "同じ内容のシートと重なっている可能性（二重に数えます。コピーなら skip）") if x)

    dst.mkdir(parents=True, exist_ok=True)
    map_path.parent.mkdir(parents=True, exist_ok=True)
    smap.to_csv(map_path, index=False, encoding="utf-8-sig")
    # 前回までの出力を片付ける（残っていると、名前や種別を変えたシートが二重に数えられる）
    manifest = dst / "_generated.txt"
    if manifest.exists():
        for old in manifest.read_text(encoding="utf-8").splitlines():
            if old.strip():
                (dst / old).unlink(missing_ok=True)
    else:
        stale = sorted(dst.glob("*.csv"))
        if stale:  # 古い版が作ったCSV: 消さずに退避する
            backup = dst.parent / f"raw_backup_{datetime.now():%Y%m%d_%H%M%S}"
            backup.mkdir(parents=True, exist_ok=True)
            for p in stale:
                shutil.move(str(p), str(backup / p.name))
            print(f"前回までの {len(stale)} 件のCSVを {backup} に移しました（二重に数えるのを防ぐため。不要なら削除してください）")
    written_names = []
    view = map_path.parent / "standardized.xlsx"
    with pd.ExcelWriter(view) as w:
        for name, parts in sorted(merged.items()):
            df = pd.concat(parts, ignore_index=True)  # 同じ公演が複数シートに分かれていたら結合
            df.to_csv(dst / f"{name}.csv", index=False, encoding="utf-8-sig")
            written_names.append(f"{name}.csv")
            title = unicodedata.normalize("NFKC", name)[:31]
            while title in w.sheets:  # Excel のシート名は31文字まで。重複したら番号を付ける
                title = title[:28] + f"~{len(w.sheets)}"
            df.to_excel(w, sheet_name=title, index=False)
        if not merged:
            pd.DataFrame().to_excel(w, sheet_name="empty")
    manifest.write_text("\n".join(written_names), encoding="utf-8")
    report_path = map_path.parent / "standardize_report.csv"
    pd.DataFrame(report, columns=["source_file", "source_sheet", "message"]).to_csv(
        report_path, index=False, encoding="utf-8-sig")

    # 個人情報は表示せず件数のみ
    todo = smap[(smap["category"] != "skip") & (smap["target_sheet"] == "")]
    guessed = smap[smap["note"].str.contains("推測|確認", na=False) & (smap["target_sheet"] != "")]
    print(f"元のシート: {len(sheets)} / 整形できたシート: {len(merged)}（{dst}/）")
    if merged:  # 列の指定が合っているか、元のシートの合計と見比べるための表示（金額の合計だけで、個人情報は出さない）
        print("シートごとの行数と金額の合計（元のシートの合計と見比べてください）:")
        for name, parts in sorted(merged.items()):
            df = pd.concat(parts, ignore_index=True)
            sums = [f"{col} {int(pd.to_numeric(df[col], errors='coerce').fillna(0).sum()):,}円"
                    for col in ("Ticket", "Support", "Entry_fee", "total_cost") if col in df.columns]
            print(f"  {name}: {len(df)} 行" + (" / " + " / ".join(sums) if sums else ""))
    if added:
        print(f"新しく見つかったシート: {len(added)} 件を {map_path} に追加しました")
    for target, n in combined:
        print(f"  {n}枚のシートを1つにまとめました: {target}")
    if dup_pairs:
        print(f"★ 同じ人・同じ金額が重なっているシートの組: {dup_pairs} 組 → {map_path} の note 列を確認してください（コピーなら skip）")
    if inherited:
        print(f"シート名を変えたシートの設定を引き継ぎました: {inherited} 件")
    if orphans:
        print(f"元のシートが見つからない行: {orphans} 件（シート名を変えたときなど）→ {map_path} の note 列を見て、不要な行は削除してください")
    if total_red:
        print(f"赤く塗られた行を除外しました: {total_red} 行（取り込む場合は --include-red）")
    if guessed_cols:
        print(f"列名の行がないシート: {guessed_cols} 件 → {map_path} の columns 列に列の推測を書きました。確認してください")
    if len(guessed):
        print(f"推測で分類したシート: {len(guessed)} 件 → {map_path} の category / date / name を確認してください")
    if len(todo):
        print(f"★ 記入が必要なシート: {len(todo)} 件 → {map_path} を開いて記入し、もう一度実行してください")
    if report:
        print(f"注意事項: {len(report)} 件 → {report_path}")
    print(f"確認用: {view}")
    return len(todo)


if __name__ == "__main__":
    main()
